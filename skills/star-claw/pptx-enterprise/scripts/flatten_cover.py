#!/usr/bin/env python3
"""Bake a template cover (its slide + the layout's background images) into a self-contained 1-slide .pptx,
filling 材料标题/部门/作者 so it can be merged into a generated deck.

usage: flatten_cover.py <template.pptx> <out.pptx> [--keep N] [--title T] [--dept D] [--author A]
"""
import io
import sys
from pptx import Presentation
from pptx.oxml.ns import qn

from template_fill import _fill_text_frame


def _parse(argv):
    pos, opts, i = [], {}, 0
    while i < len(argv):
        a = argv[i]
        if a.startswith("--"):
            opts[a[2:]] = argv[i + 1] if i + 1 < len(argv) else ""
            i += 2
        else:
            pos.append(a)
            i += 1
    return pos, opts


def main():
    pos, opts = _parse(sys.argv[1:])
    if len(pos) < 2:
        print("usage: flatten_cover.py <template.pptx> <out.pptx> [--keep N] [--title T] [--dept D] [--author A]", file=sys.stderr)
        sys.exit(2)
    template, out = pos[0], pos[1]
    keep = int(opts.get("keep", 0))
    prs = Presentation(template)
    cover = prs.slides[keep]

    # Bake the layout's background pictures onto the slide (sent to back) so the slide is self-contained.
    for shp in cover.slide_layout.shapes:
        if shp.shape_type == 13:  # PICTURE
            pic = cover.shapes.add_picture(io.BytesIO(shp.image.blob), shp.left, shp.top, shp.width, shp.height)
            el = pic._element
            el.getparent().remove(el)
            cover.shapes._spTree.insert(2, el)  # after nvGrpSpPr + grpSpPr → back of z-order

    # Fill the editable fields by matching their current text.
    for shp in cover.shapes:
        if not shp.has_text_frame:
            continue
        t = shp.text_frame.text
        if "材料标题" in t and "title" in opts:
            _fill_text_frame(shp.text_frame, opts["title"])
        elif ("部门" in t or "作者" in t) and ("dept" in opts or "author" in opts):
            _fill_text_frame(shp.text_frame, [opts.get("dept", "部门"), opts.get("author", "作者")])

    # Drop all other slides.
    lst = prs.slides._sldIdLst
    for i, el in enumerate(list(lst)):
        if i != keep:
            rId = el.get(qn("r:id"))
            try:
                prs.part.drop_rel(rId)
            except Exception:
                pass
            lst.remove(el)

    prs.save(out)
    print(f"wrote {out}")


if __name__ == "__main__":
    main()
