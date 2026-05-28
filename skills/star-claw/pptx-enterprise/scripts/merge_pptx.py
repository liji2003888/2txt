#!/usr/bin/env python3
"""Merge slides from one deck into another, copying shapes + images (charts are not copied).

usage: merge_pptx.py <base.pptx> <add.pptx> <out.pptx> [--prepend]
Intended use: prepend a flattened cover into a generated content deck.
"""
import copy
import os
import sys
import tempfile

from pptx import Presentation
from pptx.oxml.ns import qn

CHART_URI = "http://schemas.openxmlformats.org/drawingml/2006/chart"


def _blank_layout(prs):
    return min(prs.slide_layouts, key=lambda layout: len(list(layout.placeholders)))


def _add_image(slide_part, blob):
    with tempfile.NamedTemporaryFile(suffix=".img", delete=False) as f:
        f.write(blob)
        path = f.name
    try:
        _, rId = slide_part.get_or_add_image_part(path)
    finally:
        os.unlink(path)
    return rId


def _append(base, src_slide):
    new = base.slides.add_slide(_blank_layout(base))
    for ph in list(new.placeholders):
        ph._element.getparent().remove(ph._element)
    for shp in src_slide.shapes:
        new.shapes._spTree.append(copy.deepcopy(shp._element))
    for blip in list(new.shapes._spTree.iter(qn("a:blip"))):
        old = blip.get(qn("r:embed"))
        if not old:
            continue
        try:
            part = src_slide.part.related_part(old)
        except KeyError:
            continue
        blip.set(qn("r:embed"), _add_image(new.part, part.blob))
    dropped = 0
    for gf in list(new.shapes._spTree.iter(qn("p:graphicFrame"))):
        gd = gf.find(qn("a:graphic") + "/" + qn("a:graphicData"))
        if gd is not None and CHART_URI in (gd.get("uri") or ""):
            gf.getparent().remove(gf)
            dropped += 1
    if dropped:
        print(f"  warning: dropped {dropped} chart(s) on a merged slide (charts are not copied)", file=sys.stderr)
    return new


def _move(base, frm, to):
    lst = base.slides._sldIdLst
    ids = list(lst)
    el = ids[frm]
    lst.remove(el)
    lst.insert(to, el)


def main():
    flags = [a for a in sys.argv[1:] if a.startswith("--")]
    args = [a for a in sys.argv[1:] if not a.startswith("--")]
    if len(args) < 3:
        print("usage: merge_pptx.py <base.pptx> <add.pptx> <out.pptx> [--prepend]", file=sys.stderr)
        sys.exit(2)
    base = Presentation(args[0])
    add = Presentation(args[1])
    start = len(list(base.slides))
    n = len(list(add.slides))
    for s in add.slides:
        _append(base, s)
    if "--prepend" in flags:
        for i in range(n):
            _move(base, start + i, i)
    base.save(args[2])
    print(f"merged {n} slide(s) -> {args[2]} ({len(list(base.slides))} total)")


if __name__ == "__main__":
    main()
