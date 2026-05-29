#!/usr/bin/env python3
"""Extract all text (and speaker notes) from a .pptx for content QA / reading.
Prefers markitdown if installed (clean markdown); falls back to python-pptx.

usage: extract_text.py <file.pptx>
"""
import sys
from pathlib import Path


def via_markitdown(path):
    try:
        from markitdown import MarkItDown
        return MarkItDown().convert(path).text_content
    except Exception:
        return None


def via_pptx(path):
    from pptx import Presentation
    prs = Presentation(path)
    out = []
    for i, slide in enumerate(prs.slides, 1):
        out.append(f"\n## Slide {i}")
        for shape in slide.shapes:
            if shape.has_text_frame and shape.text_frame.text.strip():
                out.append(shape.text_frame.text)
            if shape.has_table:
                for row in shape.table.rows:
                    out.append(" | ".join(c.text for c in row.cells))
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame.text.strip():
            out.append(f"[notes] {slide.notes_slide.notes_text_frame.text}")
    return "\n".join(out)


def main():
    if len(sys.argv) < 2:
        print("usage: extract_text.py <file.pptx>", file=sys.stderr)
        sys.exit(2)
    path = sys.argv[1]
    text = via_markitdown(path) or via_pptx(path)
    print(text)


if __name__ == "__main__":
    main()
