#!/usr/bin/env python3
"""Dump an existing .pptx structure (slides, shapes, text, tables, images) as JSON for targeted in-place editing."""
import json
import sys
from pathlib import Path

from pptx import Presentation


def emu_px(v):
    return round(v / 9525) if v is not None else None


def shape_info(sh, idx):
    info = {"index": idx, "name": sh.name, "type": str(sh.shape_type)}
    try:
        info["pos"] = {"left": emu_px(sh.left), "top": emu_px(sh.top), "w": emu_px(sh.width), "h": emu_px(sh.height)}
    except Exception:
        pass
    try:
        if sh.is_placeholder:
            info["placeholder"] = {"idx": sh.placeholder_format.idx, "type": str(sh.placeholder_format.type)}
    except Exception:
        pass
    if sh.has_text_frame and sh.text_frame.text:
        info["text"] = sh.text_frame.text
    if sh.has_table:
        t = sh.table
        info["table"] = {
            "rows": len(t.rows),
            "cols": len(t.columns),
            "cells": [[c.text for c in row.cells] for row in t.rows],
        }
    if sh.has_chart:
        try:
            info["chart"] = {"type": str(sh.chart.chart_type)}
        except Exception:
            info["chart"] = {}
    if sh.shape_type == 13:
        info["picture"] = True
    return info


def main():
    if len(sys.argv) < 2:
        print("usage: dump_pptx.py <file.pptx> [out.json]", file=sys.stderr)
        sys.exit(2)
    prs = Presentation(sys.argv[1])
    deck = {
        "size": {"w_px": emu_px(prs.slide_width), "h_px": emu_px(prs.slide_height)},
        "slides": [],
    }
    for si, slide in enumerate(prs.slides):
        deck["slides"].append({
            "index": si,
            "layout": slide.slide_layout.name,
            "shapes": [shape_info(sh, i) for i, sh in enumerate(slide.shapes)],
        })
    out = json.dumps(deck, ensure_ascii=False, indent=2)
    if len(sys.argv) >= 3:
        Path(sys.argv[2]).write_text(out, encoding="utf-8")
        print(f"wrote {sys.argv[2]}")
    else:
        print(out)


if __name__ == "__main__":
    main()
