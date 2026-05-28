#!/usr/bin/env python3
"""Apply in-place edits to a .pptx (text/table/chart/image/slide ops) preserving the original design."""
import copy
import json
import sys
from pathlib import Path

from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.oxml.ns import qn

from template_fill import _fill_text_frame


def _find_shape(slide, op):
    if "shape_index" in op:
        return slide.shapes[op["shape_index"]]
    name = op.get("shape")
    if name:
        for sh in slide.shapes:
            if sh.name == name:
                return sh
    return None


def _replace_in_tf(tf, find, repl):
    n = 0
    for para in tf.paragraphs:
        for run in para.runs:
            if find in run.text:
                run.text = run.text.replace(find, repl)
                n += 1
        if find in para.text and find not in "".join(r.text for r in para.runs):
            # cross-run match: collapse into the first run (inline formatting within the paragraph is lost)
            joined = "".join(r.text for r in para.runs).replace(find, repl)
            if para.runs:
                para.runs[0].text = joined
                for r in para.runs[1:]:
                    r._r.getparent().remove(r._r)
                n += 1
    return n


def op_replace_text(prs, op):
    slides = [prs.slides[op["slide"]]] if "slide" in op else list(prs.slides)
    total = 0
    for slide in slides:
        for sh in slide.shapes:
            if sh.has_text_frame:
                total += _replace_in_tf(sh.text_frame, op["find"], op["replace"])
            if sh.has_table:
                for row in sh.table.rows:
                    for cell in row.cells:
                        total += _replace_in_tf(cell.text_frame, op["find"], op["replace"])
    print(f"  replace_text {op['find']!r}->{op['replace']!r}: {total} run(s)")


def op_set_text(prs, op):
    sh = _find_shape(prs.slides[op["slide"]], op)
    if sh is None or not sh.has_text_frame:
        raise ValueError(f"set_text: no text shape on slide {op['slide']} ({op.get('shape') or op.get('shape_index')})")
    _fill_text_frame(sh.text_frame, op["text"])
    print(f"  set_text slide {op['slide']}")


def op_set_table_cell(prs, op):
    sh = _find_shape(prs.slides[op["slide"]], op)
    if sh is None or not sh.has_table:
        raise ValueError(f"set_table_cell: no table on slide {op['slide']}")
    sh.table.cell(op["row"], op["col"]).text = str(op["text"])
    print(f"  set_table_cell slide {op['slide']} ({op['row']},{op['col']})")


def op_set_chart_data(prs, op):
    sh = _find_shape(prs.slides[op["slide"]], op)
    if sh is None or not sh.has_chart:
        raise ValueError(f"set_chart_data: no chart on slide {op['slide']}")
    cd = CategoryChartData()
    cd.categories = op["categories"]
    for s in op["series"]:
        cd.add_series(s["name"], s["values"])
    sh.chart.replace_data(cd)
    print(f"  set_chart_data slide {op['slide']}")


def op_replace_image(prs, op):
    slide = prs.slides[op["slide"]]
    sh = _find_shape(slide, op)
    if sh is None or sh.shape_type != 13:
        raise ValueError(f"replace_image: no picture on slide {op['slide']}")
    blip = sh._element.find(".//" + qn("a:blip"))
    rId = blip.get(qn("r:embed"))
    part = slide.part.related_part(rId)
    part._blob = Path(op["path"]).read_bytes()
    print(f"  replace_image slide {op['slide']} <- {op['path']}")


def _move(prs, old_index, new_index):
    lst = prs.slides._sldIdLst
    ids = list(lst)
    el = ids[old_index]
    lst.remove(el)
    lst.insert(new_index, el)


def op_duplicate_slide(prs, op):
    src = prs.slides[op["index"]]
    new = prs.slides.add_slide(src.slide_layout)
    for sh in list(new.shapes):
        sh._element.getparent().remove(sh._element)
    for sh in src.shapes:
        new.shapes._spTree.append(copy.deepcopy(sh._element))
    for rId, rel in src.part.rels.items():
        if "notesSlide" in rel.reltype or rId in new.part.rels:
            continue
        new.part.rels.add_relationship(rel.reltype, rel._target, rId, is_external=rel.is_external)
    if "to" in op:
        _move(prs, len(prs.slides) - 1, op["to"])
    print(f"  duplicate_slide {op['index']} -> {op.get('to', 'end')}")


def op_delete_slide(prs, op):
    lst = prs.slides._sldIdLst
    ids = list(lst)
    el = ids[op["index"]]
    rId = el.get(qn("r:id"))
    prs.part.drop_rel(rId)
    lst.remove(el)
    print(f"  delete_slide {op['index']}")


def op_reorder_slides(prs, op):
    lst = prs.slides._sldIdLst
    ids = list(lst)
    for el in ids:
        lst.remove(el)
    for i in op["order"]:
        lst.append(ids[i])
    print(f"  reorder_slides {op['order']}")


OPS = {
    "replace_text": op_replace_text,
    "set_text": op_set_text,
    "set_table_cell": op_set_table_cell,
    "set_chart_data": op_set_chart_data,
    "replace_image": op_replace_image,
    "duplicate_slide": op_duplicate_slide,
    "delete_slide": op_delete_slide,
    "reorder_slides": op_reorder_slides,
}


def main():
    if len(sys.argv) < 4:
        print("usage: edit_pptx.py <in.pptx> <ops.json> <out.pptx>", file=sys.stderr)
        print("  ops.json: {\"ops\": [{\"op\": \"replace_text\", \"find\": ..., \"replace\": ...}, ...]}", file=sys.stderr)
        sys.exit(2)
    prs = Presentation(sys.argv[1])
    spec = json.loads(Path(sys.argv[2]).read_text(encoding="utf-8"))
    for op in spec.get("ops", []):
        fn = OPS.get(op.get("op"))
        if fn is None:
            raise ValueError(f"unknown op: {op.get('op')}")
        fn(prs, op)
    prs.save(sys.argv[3])
    print(f"wrote {sys.argv[3]}")


if __name__ == "__main__":
    main()
